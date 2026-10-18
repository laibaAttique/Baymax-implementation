"""Builds Baymax_Chapter5_Complete_Testing_Modules_4_5_6.docx from the original Chapter 5 and the recorded test results.

Nothing in the original is deleted, reworded or renumbered. New sections, tables and figures are inserted (cloned from the original's
own elements so formatting matches); the only edits to existing content are the aggregate counts in Table 5.1 and Table 5.31.
Every Actual Result / Status comes from a results file written by a real run:
  results/atc_m456.jsonl            Flutter ATC cases (ATC_LOG output of test/atc/atc_m456_*_test.dart)
  ../../automation/appium/reports/  Appium run (JUnit XML of the failed session start; appium server log)
Run:  python tools/build_chapter5.py   (from docs/chapter5, with python-docx and Pillow)
"""
import copy
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent.parent  # docs/chapter5
REPO = HERE.parent.parent
ORIGINAL = HERE / 'Baymax_Chapter5_Test_Cases.original.docx'
OUT = HERE / 'Baymax_Chapter5_Complete_Testing_Modules_4_5_6.docx'
RESULTS = HERE / 'results' / 'atc_m456.jsonl'
APPIUM = REPO / 'automation' / 'appium'
FIG_DIR = HERE / 'figures'
FR_M56 = HERE / 'tools' / 'fr_m56.json'

PASS, FAIL, BLOCKED = '1B7F3B', 'C00000', '9C5700'

# ---------------------------------------------------------------- data
cases = OrderedDict()
for line in RESULTS.read_text(encoding='utf-8').splitlines():
    d = json.loads(line)
    cases[d['id']] = d
assert len(cases) == 123, len(cases)
assert all(c['ok'] for c in cases.values()), [c['id'] for c in cases.values() if not c['ok']]

# Appium cases: metadata from the spec files; status from the real run (no session could be created).
junit = sorted((APPIUM / 'reports' / 'junit').glob('*.xml'))
junit_text = '\n'.join(p.read_text(encoding='utf-8') for p in junit)
assert 'WRITE_SECURE_SETTINGS' in junit_text and 'hidden_api_policy' in junit_text, 'expected the recorded session failure'
APPIUM_ACTUAL = ('Not executed. Appium could not create a session on the Redmi Note 11: during its standard start-up the UiAutomator2 '
                 'driver runs "settings put global hidden_api_policy 1", and MIUI refused it (SecurityException: requires '
                 'WRITE_SECURE_SETTINGS; MIUI "USB debugging (Security settings)" is off). Taps are refused for the same reason '
                 '(adb input tap: "Injecting input events requires ... INJECT_EVENTS"). Evidence: automation/appium/reports/.')
appium = []
for spec in sorted((APPIUM / 'tests').glob('*.e2e.js')):
    src = spec.read_text(encoding='utf-8')
    for m in re.finditer(r"await atc\(this, \{\s*id: '([^']+)',\s*req: '([^']+)',\s*objective: (['`])(.*?)\3,\s*input: (['`])(.*?)\5,\s*expected: (['`])(.*?)\7,", src, re.S):
        exp = m.group(8).replace('${PACKAGE}', 'com.example.app')
        appium.append(dict(id=m.group(1), req=m.group(2), objective=m.group(4).replace("\\'", "'"), input=m.group(6).replace("\\'", "'"),
                           expected=exp.replace("\\'", "'"), actual=APPIUM_ACTUAL, status='Blocked', spec=spec.name,
                           tool='Appium 3.8.0 + UiAutomator2 8.7.0 (WebdriverIO 9.32.0), real device'))
assert len(appium) == 12, len(appium)
with open(APPIUM / 'reports' / 'appium_results.jsonl', 'w', encoding='utf-8') as f:
    for a in appium:
        f.write(json.dumps(a) + '\n')

fr = json.loads(FR_M56.read_text(encoding='utf-8'))
sys.path.insert(0, str(REPO / 'docs' / 'test-cases' / 'tools'))
import requirements  # noqa: E402  (team file: M4 FR texts)
m4 = {f'M4-FR{n}': t for n, t, _ in requirements.REQS[4]}
SFR = {
    'M4-SFR1': 'Import the free-exercise-db exercise dataset with stable, unique IDs (added during implementation)',
    'M4-SFR2': 'Match each exercise to a Compendium MET value; estimate energy as MET x weight x time (added)',
    'M4-SFR3': 'Screen readiness (warning symptoms, clearance, missing data) and owned equipment before planning (added)',
    'M4-SFR4': 'Record completion of planned exercises from sessions and mark the plan day done (added)',
    'M5-SFR1': 'Map every exercise to a movement pattern with a camera tracker or a stated reason (added)',
    'M5-SFR2': 'Time static holds (plank, wall sit) in camera mode (added)',
    'M5-SFR3': 'Start posture sessions from the personalised plan (Today\'s workout / My plan) as well as the library (added)',
    'M6-SFR1': 'Recognise workout intents from the spoken (transcribed) text (added)',
    'M6-SFR2': 'Speak posture corrections and answer form questions from the live session (added)',
    'M6-SFR3': 'Control the posture session by voice (pause, resume, stop with confirmation) (added)',
}
def req_text(r):
    if r in m4:
        return m4[r]
    if r in SFR:
        return SFR[r]
    return fr[r.split('-', 1)[1]]

# Cases that also verify further requirements (they check those labels/values explicitly), used only for the RTM.
ALSO = {
    'ATC-M5-F09': ['M5-FR18.1', 'M5-FR18.2', 'M5-FR18.3'], 'ATC-M5-F10': ['M5-FR17.3', 'M5-FR17.4'],
    'ATC-M5-F05': ['M5-FR15.1', 'M5-FR15.2'], 'ATC-M5-F07': ['M5-FR15.2'], 'ATC-M5-F08': ['M5-FR16.4', 'M5-FR16.6'],
    'ATC-M6-F01': ['M6-FR19.2'], 'ATC-M6-F02': ['M6-FR19.5'], 'ATC-M5-A02': ['M5-FR15.2', 'M5-FR15.3', 'M5-FR15.4'],
    'ATC-M5-A04': ['M5-FR17.3', 'M5-FR17.4'], 'ATC-M6-A01': ['M6-FR19.2'],
}

def group(prefixes):
    return [c for i, c in cases.items() if any(i.startswith(p) for p in prefixes)]

# ---------------------------------------------------------------- document helpers
d = docx.Document(str(ORIGINAL))
body = d.element.body
orig = list(body.iterchildren())
assert len(orig) == 241

T = lambda i: orig[i]
TPL = dict(h3=T(107), caption=T(108), normal=T(18), spacer=T(20), unit_tbl=T(19), func_tbl=T(64), int_tbl=T(201), dt=T(109),
           br_tbl=T(112), tools=T(220), cases=T(224), cmds=T(229), results=T(232), rtm=T(238), fig_caption=T(234), h3_55=T(236))

def clone(el):
    return copy.deepcopy(el)

def set_par_text(p_el, text):
    runs = p_el.findall(qn('w:r'))
    for r in runs[1:]:
        p_el.remove(r)
    t = runs[0].find(qn('w:t'))
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return p_el

def par(kind, text):
    return set_par_text(clone(TPL[kind]), text)

def spacer():
    return clone(TPL['spacer'])

def set_cell(tc, text, color=None):
    ps = tc.findall(qn('w:p'))
    for p in ps[1:]:
        tc.remove(p)
    p = ps[0]
    runs = p.findall(qn('w:r'))
    for r in runs[1:]:
        p.remove(r)
    r = runs[0]
    t = r.find(qn('w:t'))
    t.text = str(text)
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    if color:
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = r.makeelement(qn('w:rPr'), {})
            r.insert(0, rpr)
        c = rpr.find(qn('w:color'))
        if c is None:
            c = rpr.makeelement(qn('w:color'), {})
            later = [e for e in rpr if e.tag in (qn('w:sz'), qn('w:szCs'), qn('w:highlight'), qn('w:u'), qn('w:lang'))]
            if later:
                later[0].addprevious(c)
            else:
                rpr.append(c)
        c.set(qn('w:val'), color)

def table_from(template, rows, header=None, status_col=None):
    tbl = clone(template)
    trs = tbl.findall(qn('w:tr'))
    head, data = trs[0], trs[1]
    for tr in trs[1:]:
        tbl.remove(tr)
    if header:
        for tc, h in zip(head.findall(qn('w:tc')), header):
            set_cell(tc, h)
    for values in rows:
        tr = clone(data)
        for k, (tc, v) in enumerate(zip(tr.findall(qn('w:tc')), values)):
            color = None
            if status_col is not None and k == status_col:
                color = PASS if str(v).startswith('Pass') else BLOCKED if str(v).startswith(('Blocked', 'Not')) else FAIL
            set_cell(tc, v, color)
        tbl.append(tr)
    return tbl

def decision_table(conds_order, rules, actions):
    """conds_order: condition names; rules: list of dicts; actions: list of action strings (one per rule)."""
    tbl = clone(TPL['dt'])
    n = len(rules)
    total, first = 14394, 3000
    w = (total - first) // n
    grid = tbl.find(qn('w:tblGrid'))
    for g in list(grid):
        grid.remove(g)
    for k in range(n + 1):
        g = grid.makeelement(qn('w:gridCol'), {qn('w:w'): str(first if k == 0 else w)})
        grid.append(g)
    trs = tbl.findall(qn('w:tr'))
    head, cond_row, act_row = trs[0], trs[1], trs[-1]
    for tr in trs:
        tbl.remove(tr)

    def build(template, label, values):
        tr = clone(template)
        tcs = tr.findall(qn('w:tc'))
        first_tc, other = tcs[0], tcs[1]
        for tc in tcs:
            tr.remove(tc)
        a = clone(first_tc)
        set_cell(a, label)
        tr.append(a)
        for v in values:
            c = clone(other)
            c.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(w))
            set_cell(c, v)
            tr.append(c)
        return tr
    tbl.append(build(head, 'Conditions', [f'Rule {k + 1}' for k in range(n)]))
    for name in conds_order:
        tbl.append(build(cond_row, name, [r.get(name, '-') for r in rules]))
    tbl.append(build(act_row, 'Action / expected result', actions))
    return tbl

def result(c):
    return 'Pass' if c['ok'] else 'Fail'

def test_rows(cs):
    return [[k + 1, c['id'], c['objective'], c['input'], c['expected'], c['actual'], result(c)] for k, c in enumerate(cs)]

def appium_rows(cs, start=1):
    return [[start + k, a['id'], a['objective'], a['input'], a['expected'], a['actual'], 'Blocked'] for k, a in enumerate(cs)]

def insert_before(anchor, elements):
    for e in elements:
        anchor.addprevious(e)

def insert_after(anchor, elements):
    for e in reversed(elements):
        anchor.addnext(e)

table_no = [32]
def caption(text):
    table_no[0] += 1
    return par('caption', f'Table 5.{table_no[0]}  {text}')

# ---------------------------------------------------------------- 5.1 Unit testing (before "5.2 Functional Testing")
unit = [
    ('Unit Testing 12: Exercise knowledge base, safety rules and plan generator (Module 4, Workout & Exercise)',
     'Testing Objective: To ensure the exercise dataset is imported faithfully, the safety and screening rules decide each exercise as written, '
     'MET-based energy estimates are correct, and the weekly plan respects the training profile of each health category.', group(['ATC-M4-U'])),
    ('Unit Testing 13: Posture engine, repetition counting and form checks (Module 5, Posture Detection)',
     'Testing Objective: To ensure joint angles, landmark visibility, movement phases, repetition acceptance, hold timing, form corrections and '
     'the posture score are computed correctly from pose landmarks. Inputs are synthetic landmark frames: these cases verify the logic, '
     'not detection accuracy on real people.', group(['ATC-M5-U'])),
    ('Unit Testing 14: Intent recognition, coach replies and voice session control (Module 6, Voice Assistant)',
     'Testing Objective: To ensure spoken (transcribed) requests map to the right intent, replies use the live session data instead of invented '
     'numbers, and microphone, recognition and speech-output failures are handled. A simulated recogniser and speech output are used: these '
     'cases do not validate microphone capture or speech-recognition accuracy.', group(['ATC-M6-U'])),
]
els = []
for title, objective, cs in unit:
    els += [par('h3', title), par('normal', objective), table_from(TPL['unit_tbl'], test_rows(cs), status_col=6), spacer()]
insert_before(T(49), els)

# ---------------------------------------------------------------- 5.2 Functional testing (before "5.3 Business Rules Testing")
appium_by = {a['id']: a for a in appium}
func = [
    ('Functional Testing 15: Exercise detail and workout plan screens (Module 4, Workout & Exercise)', group(['ATC-M4-F']), []),
    ('Functional Testing 16: Posture detection home, mode selection, camera, manual guidance and summary screens (Module 5, Posture Detection)',
     group(['ATC-M5-F']), []),
    ('Functional Testing 17: Voice assistant screen (Module 6, Voice Assistant)', group(['ATC-M6-F']), []),
]
els = []
for title, cs, _ in func:
    els += [par('h3', title), par('normal', "Objective: to ensure the screen shows the correct content and reacts correctly to the user's actions."),
            table_from(TPL['func_tbl'], test_rows(cs), status_col=6), spacer()]
dev_func = [a for a in appium if not a['id'].startswith('ATC-INT')]
els += [par('h3', 'Functional Testing 18: Real-device UI automation of Modules 4, 5 and 6 (Appium + UiAutomator2, Redmi Note 11)'),
        par('normal', 'Objective: to run the same screens on the real phone through the Android accessibility tree (labels, never screen '
                      'coordinates). These cases were implemented and run, but the phone refused the Appium session, so none of them executed; '
                      'they are reported as Blocked, not as passed. UI automation alone would not prove posture-detection or speech-recognition accuracy.'),
        table_from(TPL['func_tbl'], appium_rows(dev_func), status_col=6), spacer()]
insert_before(T(106), els)

# ---------------------------------------------------------------- 5.3 Business rules (before "5.4 Integration Testing")
dts = OrderedDict()
for c in cases.values():
    if 'table' in c and re.match(r'ATC-M[456]-B', c['id']):
        dts.setdefault(c['table'], []).append(c)
els = []
br_no = 13
for table, cs in dts.items():
    br_no += 1
    code = table.split(' ')[0]
    conds = list(OrderedDict.fromkeys(k for c in cs for k in c['conds']))
    els += [par('h3', f'Business Rule Testing {br_no}: {table}'), caption(f'Decision table: {table}'),
            decision_table(conds, [c['conds'] for c in cs], [c['actions'] for c in cs]), spacer(), caption(f'Test cases for {code}')]
    rows = [[k + 1, c['id'], f'Rule {k + 1}: {c["objective"]}', c['input'], c['expected'], c['actual'], result(c)] for k, c in enumerate(cs)]
    els += [table_from(TPL['br_tbl'], rows, status_col=6), spacer()]
insert_before(T(198), els)
assert table_no[0] == 32 + 2 * len(dts)

# ---------------------------------------------------------------- 5.4 Integration (before "5.5 Automated Testing")
integ = [
    ('Integration Testing 5: Module 2 health classification to Module 4 workout plan', ['ATC-INT-34', 'ATC-INT-35'],
     'Testing Objective: to ensure the health category and conditions from Module 2 set the restrictions and intensity of the generated plan.'),
    ('Integration Testing 6: Module 4 workout plan and Module 5 posture sessions (both directions)', ['ATC-INT-36', 'ATC-INT-37', 'ATC-INT-38'],
     'Testing Objective: to ensure a planned exercise starts a posture session with its planned dose, and that only a session that reached its '
     'target completes the planned exercise and the plan day.'),
    ('Integration Testing 7: Module 5 posture detection and Module 6 voice assistant (both directions)',
     ['ATC-INT-39', 'ATC-INT-40', 'ATC-INT-41', 'ATC-INT-42', 'ATC-INT-43'],
     'Testing Objective: to ensure live posture results reach the voice coach (spoken corrections, form and rep answers) and voice commands '
     'control the posture session. Pose frames, microphone and speaker are simulated.'),
    ('Integration Testing 8: End-to-end Modules 2, 4, 5 and 6', ['ATC-INT-44'],
     'Testing Objective: to ensure a profile produces a plan whose exercises can be done as sessions, queried by voice and recorded as a completed plan day.'),
]
els = []
for title, ids, objective in integ:
    els += [par('h3', title), par('normal', objective), table_from(TPL['int_tbl'], test_rows([cases[i] for i in ids]), status_col=6), spacer()]
dev_int = [a for a in appium if a['id'].startswith('ATC-INT')]
els += [par('h3', 'Integration Testing 9: Cross-module flows on the real phone (Appium + UiAutomator2)'),
        par('normal', 'Testing Objective: to follow Module 4 to Module 5 and Module 5 to Module 6 navigation on the real phone. Implemented and run; '
                      'blocked by the phone before any step executed.'),
        table_from(TPL['int_tbl'], appium_rows(dev_int), status_col=6), spacer()]
insert_before(T(215), els)

# ---------------------------------------------------------------- 5.5 Automated testing
insert_after(T(217), [par('normal',
    'For Modules 4 (Workout Plan Generation), 5 (Posture Detection) and 6 (Voice Assistant), 123 further cases were added to the same Flutter '
    'ATC suites (test/atc/atc_m456_*_test.dart) and recorded by the same ATC_LOG mechanism, so their Actual Result columns are the output of a '
    'real run on 2026-10-18. A real-device layer was added with Appium and the UiAutomator2 driver driven by WebdriverIO (automation/appium/); '
    'its selectors use the accessibility labels Flutter exposes, never screen coordinates. On the test phone those cases were blocked by the '
    'device (see 5.5.4) and are reported as Blocked. Module 6 is the voice assistant; the Module 13 AI chatbot is only its entry point.')])

tools_rows = [
    ['flutter_test (Flutter test framework)', 'Unit, widget (screen), decision-table and cross-module cases of Modules 4, 5 and 6',
     'Flutter 3.47.7, Dart 3.13.5 (stable)', 'Unit / Functional / Business rules / Integration', 'Windows 11 Home, headless (no phone needed)'],
    ['Appium server', 'Real-device UI automation server (W3C WebDriver)', 'appium 3.8.0', 'Functional / Integration (real device)', 'Node.js v24.14.1 on Windows 11'],
    ['Appium UiAutomator2 driver', 'Drives the Android accessibility tree of the installed app', 'appium-uiautomator2-driver 8.7.0', 'Functional / Integration (real device)', 'Redmi Note 11 (2201117TG), Android 13, MIUI'],
    ['WebdriverIO (+ Mocha framework, spec and JUnit reporters, appium service)', 'Test runner and JUnit XML output for the Appium cases',
     'webdriverio / @wdio/* 9.32.0', 'Functional / Integration (real device)', 'Node.js v24.14.1, npm 11.11.0'],
    ['Android Debug Bridge (adb)', 'Device connection, app launch, screenshots and accessibility dump', 'adb 1.0.41 (platform-tools 37.0.1)', '-', 'Windows 11, USB to the phone'],
    ['Java runtime', 'Needed by the UiAutomator2 driver tooling', 'JDK 21.0.10 (JAVA_HOME)', '-', 'Windows 11'],
]
insert_before(T(222), [caption('Automated testing tools added for Modules 4, 5 and 6'), table_from(TPL['tools'], tools_rows), spacer()])

all_rows = []
for c in cases.values():
    all_rows.append([c['id'], c['req'], c['objective'], c['input'], c['expected'], c['actual'], result(c)])
for a in appium:
    all_rows.append([a['id'], a['req'], a['objective'], a['input'], a['expected'], a['actual'], 'Blocked'])
insert_before(T(226), [caption('Automated test cases added for Modules 4, 5 and 6'), table_from(TPL['cases'], all_rows, status_col=6), spacer()])

cmd_rows = [
    ['Modules 4-6 unit cases', 'cd app && flutter test test/atc/atc_m456_unit_test.dart', 'Unit'],
    ['Modules 4-6 functional cases', 'cd app && flutter test test/atc/atc_m456_functional_test.dart', 'Functional'],
    ['Modules 4-6 business-rule cases', 'cd app && flutter test test/atc/atc_m456_rules_test.dart', 'Business rules'],
    ['Modules 4-6 integration cases', 'cd app && flutter test test/atc/atc_m456_integration_test.dart', 'Integration'],
    ['Real-device cases (Appium)', 'cd automation/appium && npm install && npx wdio run ./wdio.conf.js', 'Functional / Integration (real device)'],
    ['Complete Flutter regression run', 'cd app && flutter test --concurrency=2', 'All app tests'],
]
insert_before(T(231), [caption('How the Modules 4, 5 and 6 suites are run'), table_from(TPL['cmds'], cmd_rows), spacer(),
                       par('normal', 'Set ATC_LOG=<file> to record each Flutter case as one JSON line (results/atc_m456.jsonl); the Appium run writes '
                                     'JUnit XML to automation/appium/reports/junit and one JSON line per case to reports/appium_results.jsonl.'), spacer()])

# Figures from the real runs.
FIG_DIR.mkdir(exist_ok=True)
log = (HERE / 'results' / 'flutter_m456_run.txt').read_text(encoding='utf-8').splitlines()
log = [l.replace('D:/fyp/Baymax-Phase 2/Baymax-implementation/app/', '') for l in log if re.match(r'\d\d:\d\d \+', l)]
shown = ['> flutter test test/atc/atc_m456_unit_test.dart ..._functional_test.dart ..._rules_test.dart ..._integration_test.dart',
         '  (2026-10-18, Windows 11, Flutter 3.47.7; test-progress lines of the console output; debug prints omitted)', ''] + [l[:150] for l in log[-24:]]
font = ImageFont.truetype('C:/Windows/Fonts/consola.ttf', 18)
img = Image.new('RGB', (1500, 28 * len(shown) + 30), (30, 30, 30))
dr = ImageDraw.Draw(img)
for k, l in enumerate(shown):
    dr.text((16, 14 + 28 * k), l, font=font, fill=(120, 220, 140) if 'All tests passed' in l else (230, 230, 230))
img.save(FIG_DIR / 'figure_5_2_flutter_m456_run.png')
phone = APPIUM / 'screenshots' / 'device_smoke_dashboard.png'

def picture(path, width):
    p = d.add_paragraph()
    p.alignment = 1
    p.add_run().add_picture(str(path), width=width)
    el = p._p
    el.getparent().remove(el)
    return el

counts = OrderedDict((k, sum(1 for i in cases if re.match(p, i))) for k, p in
                     [('unit', r'ATC-M[456]-U'), ('functional', r'ATC-M[456]-F'), ('rules', r'ATC-M[456]-B'), ('integration', r'ATC-INT')])
res_rows = [
    ['Flutter test (Dart): Modules 4-6 unit', counts['unit'], counts['unit'], 0, 0],
    ['Flutter test (Dart): Modules 4-6 functional (widget)', counts['functional'], counts['functional'], 0, 0],
    ['Flutter test (Dart): Modules 4-6 business rules', counts['rules'], counts['rules'], 0, 0],
    ['Flutter test (Dart): Modules 4-6 integration', counts['integration'], counts['integration'], 0, 0],
    ['Appium + UiAutomator2 on the Redmi Note 11 (real device)', 0, 0, 0, '12 (Blocked by the device)'],
    ['Total for Modules 4, 5 and 6', len(cases), len(cases), 0, 12],
    ['Complete Flutter regression run, all modules (not counted above)', 993, 993, 0, '13 (skipped in source)'],
]
insert_before(T(236), [
    caption('Execution results for Modules 4, 5 and 6 (run of 2026-10-18)'), table_from(TPL['results'], res_rows), spacer(),
    par('normal', 'Appium: the WebdriverIO run reached the phone and started the UiAutomator2 driver, which stopped at its first device step '
                  '("settings put global hidden_api_policy 1"): MIUI refused it with "Permission denial: writing to settings requires '
                  'android.permission.WRITE_SECURE_SETTINGS". No setting was changed and no Appium app was left on the phone. Tapping is refused for '
                  'the same reason ("Injecting input events requires ... INJECT_EVENTS"). Both are controlled by MIUI Developer options, "USB '
                  'debugging (Security settings)", which only the phone owner can enable; the suite is ready to re-run unchanged when it is. '
                  'The regression run\'s 13 skipped tests are marked skip in the source (4 superseded guidance-screen tests, 9 food-list tests '
                  'pending licence checks).'),
    picture(FIG_DIR / 'figure_5_2_flutter_m456_run.png', Inches(8.5)), par('fig_caption', 'Figure 5.2: Execution of the Modules 4, 5 and 6 Flutter test cases (console output)'),
    picture(phone, Inches(2.3)), par('fig_caption', 'Figure 5.3: Baymax running on the Redmi Note 11: dashboard with the Module 4 "Today\'s workout" card '
                                                    '(adb screenshot; device smoke check, not an Appium result)'), spacer(),
])

# RTM additions (after Table 5.32).
order = (['M4-FR1', 'M4-FR2', 'M4-FR5', 'M4-FR6', 'M4-SFR1', 'M4-SFR2', 'M4-SFR3', 'M4-SFR4'] +
         sorted([f'M5-{k}' for k in fr if k.startswith(('FR15', 'FR16', 'FR17', 'FR18'))], key=lambda s: [int(x) for x in s[5:].split('.')]) +
         ['M5-SFR1', 'M5-SFR2', 'M5-SFR3'] + sorted([f'M6-{k}' for k in fr if k.startswith('FR19')]) + ['M6-SFR1', 'M6-SFR2', 'M6-SFR3'])
feature = {'M4': 'Workout & Exercise (plan generation)', 'M5': 'Posture Detection', 'M6': 'Voice Assistant'}
rtm_rows = []
for r in order:
    ids = [c['id'] for c in cases.values() if c['req'] == r or r in ALSO.get(c['id'], [])]
    dev = [a['id'] for a in appium if a['req'] == r or r in ALSO.get(a['id'], [])]
    if r == 'M5-FR18.4':
        rtm_rows.append([r, req_text(r), feature['M5'], '-', '-', 'No', 'Not implemented in this release (no test)'])
        continue
    if not ids and not dev:
        rtm_rows.append([r, req_text(r), feature[r[:2]], '-', '-', 'No', 'Not covered by an automated case'])
        continue
    ok = sum(1 for i in ids if cases[i]['ok'])
    status = (f'Pass ({ok}/{len(ids)})' if ids else 'Not run') + (f'; real device Blocked ({len(dev)})' if dev else '')
    tool = 'Flutter test' + ('; Appium + UiAutomator2' if dev else '') if ids else 'Appium + UiAutomator2'
    rtm_rows.append([r, req_text(r), feature[r[:2]], ', '.join(sorted(ids) + dev), tool, 'Yes', status])
insert_before(T(240), [
    caption('Requirement Traceability Matrix for Modules 4, 5 and 6'), table_from(TPL['rtm'], rtm_rows, status_col=6), spacer(),
    par('h3_55', '5.5.6 Scope and limitations of the Modules 4, 5 and 6 evidence'),
    par('normal', 'Posture detection: the engine cases use synthetic landmark frames, and the real-photo replay test uses landmarks extracted '
                  'from public-domain exercise photos (held-out half: 66% of repetitions counted, 13% false alarms). They verify the counting, timing '
                  'and feedback logic. They do not show detection accuracy on live people; the only live device check so far is the squat '
                  '(2026-10-15, docs/MODULE5_DEVICE_VALIDATION.md). Only four exercises have validated form scoring; the rest are labelled accordingly.'),
    par('normal', 'Voice assistant: recognition and speech output are simulated in every Flutter case. The cases show what the assistant does '
                  'with a transcript and with recognition or speech failures; they are not microphone or recognition-accuracy validation.'),
    par('normal', 'Workout plan: passing rule cases show that the screening and safety rules are applied exactly as written (rule set pending '
                  'professional review). They are not evidence of clinical safety, and Baymax does not give medical clearance.'),
    par('normal', 'Real device: the 12 Appium cases are implemented but Blocked by the phone\'s security setting, so no claim is made from them. '
                  'M5-FR18.4 (injury-risk notice) is not implemented and has no test. Requirement IDs marked "added" (SFR) were added during '
                  'implementation and are not yet in Chapter 2.'),
])

# ---------------------------------------------------------------- aggregate counts (the only edits to existing content)
def cell(tbl_el, r, c):
    return Table(tbl_el, d).rows[r].cells[c]

def set_count(tbl_el, r, c, old, new):
    p = cell(tbl_el, r, c).paragraphs[0]
    assert p.text == str(old), (r, c, p.text, old)
    set_par_text(p._p, str(new))

n_func_dev, n_int_dev = len(dev_func), len(dev_int)
t51 = T(2)
upd = {1: (105, 105 + counts['unit'], 0), 2: (73, 73 + counts['functional'] + n_func_dev, n_func_dev), 3: (84, 84 + counts['rules'], 0),
       4: (27, 27 + counts['integration'] + n_int_dev, n_int_dev)}
for r, (old, new, notrun) in upd.items():
    set_count(t51, r, 2, old, new)
    set_count(t51, r, 3, old, new - notrun)
    set_count(t51, r, 5, 0, notrun)
set_count(t51, 5, 2, 289, 289 + len(cases) + len(appium))
set_count(t51, 5, 3, 289, 289 + len(cases))
set_count(t51, 5, 5, 0, len(appium))
insert_after(T(2), [par('normal', f'Counts updated for Modules 4, 5 and 6: {len(cases)} Flutter cases added and passed ({counts["unit"]} unit, '
                                  f'{counts["functional"]} functional, {counts["rules"]} business-rule, {counts["integration"]} integration), and '
                                  f'{len(appium)} real-device Appium cases ({n_func_dev} functional, {n_int_dev} integration) that were blocked by the '
                                  'phone and are counted as Not run.')])
t531 = T(232)
set_count(t531, 1, 1, 203, 203 + len(cases))
set_count(t531, 1, 2, 203, 203 + len(cases))
set_count(t531, 4, 1, 289, 289 + len(cases))
set_count(t531, 4, 2, 289, 289 + len(cases))

d.core_properties.revision = (d.core_properties.revision or 1) + 1
d.save(str(OUT))
print('wrote', OUT, '| new tables up to 5.%d | decision tables %d' % (table_no[0], len(dts)))
